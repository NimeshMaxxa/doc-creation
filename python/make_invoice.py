#!/usr/bin/env python3
"""Fire Fighting Solutions (Pvt) Ltd - invoice generator.

Fill in an invoice workbook (copy invoice_template.xlsx), then run:

    python make_invoice.py my_invoice.xlsx

The PDF is written to the output/ folder as "INVOICE <number>.pdf".

Other commands:
    python make_invoice.py --template new.xlsx   write a fresh blank workbook
    python make_invoice.py --next                show the next invoice number
"""
import argparse
import datetime as dt
import os
import sys
from decimal import Decimal, ROUND_HALF_UP
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUTPUT = os.path.join(HERE, "output")
COUNTER = os.path.join(ASSETS, "last_invoice_no.txt")

# ---------------------------------------------------------------- company
COMPANY = {
    "address": "No.366, Heentatigala Road, Thalpe,Galle.",
    "email": "firefightingsolutionpro@gmail.com",
    "hotline": "0771521009",
    "tel": "0711077979",
    "bank": [
        ("Name of the Account", "FIRE FIGHTING SOLUTIONS (PVT) LTD"),
        ("Name of the Bank", "HATTON NATIONAL BANK"),
        ("Bank Account No", "237010003226"),
        ("Branch Name", "KOGGALA"),
    ],
}

# ---------------------------------------------------------------- fonts
def register_fonts():
    """Use real Calibri when installed (Windows), else the bundled Carlito,
    which has identical letter widths."""
    candidates = [
        (r"C:\Windows\Fonts\calibri.ttf", r"C:\Windows\Fonts\calibrib.ttf"),
        ("/Applications/Microsoft Word.app/Contents/Resources/DFonts/Calibri.ttf",
         "/Applications/Microsoft Word.app/Contents/Resources/DFonts/Calibrib.ttf"),
        (os.path.join(ASSETS, "fonts", "Carlito-Regular.ttf"),
         os.path.join(ASSETS, "fonts", "Carlito-Bold.ttf")),
    ]
    for regular, bold in candidates:
        if os.path.exists(regular) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("Body", regular))
            pdfmetrics.registerFont(TTFont("Body-Bold", bold))
            pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-Bold")
            return
    sys.exit("No Calibri/Carlito font found in assets/fonts.")


# ---------------------------------------------------------------- numbers
def money(value):
    d = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{d:,.2f}"


def qty_text(value):
    if isinstance(value, (int, float)) and float(value).is_integer():
        return f"{int(value):02d}"
    return str(value)


def to_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    try:
        return Decimal(str(value).replace(",", "").strip())
    except Exception:
        sys.exit(f"Not a number: {value!r}")


def read_counter():
    try:
        with open(COUNTER) as f:
            return int(f.read().strip())
    except FileNotFoundError:
        return 0


def write_counter(n):
    with open(COUNTER, "w") as f:
        f.write(f"{n}\n")


# ---------------------------------------------------------------- workbook
FIELDS = ["Invoice No", "Date", "PO No", "Quote No",
          "Customer", "Address line 1", "Address line 2", "Address line 3"]
COLUMNS = ["Item", "Description", "Remark", "Unit", "Qty", "Rate", "Amount"]


def write_template(path, sample=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="FDE9E7")

    ws["A1"] = "FFS invoice - fill the yellow cells, then run make_invoice.py"
    ws["A1"].font = Font(bold=True, size=13)
    notes = {
        "Invoice No": "leave blank = next number automatically",
        "Date": "leave blank = today",
    }
    sample = sample or {}
    for i, name in enumerate(FIELDS):
        r = 3 + i
        ws.cell(r, 1, name).font = bold
        c = ws.cell(r, 2, sample.get(name))
        c.fill = PatternFill("solid", fgColor="FFF8C5")
        c.border = box
        if name in notes:
            ws.cell(r, 4, notes[name]).font = Font(italic=True, color="777777")
        if name == "Date":
            c.number_format = "DD/MM/YYYY"

    hr = 3 + len(FIELDS) + 1
    ws.cell(hr - 1, 1, "Items  (a row with only a Description = bold section "
            "heading; Amount only = lump sum; blank Item = auto number; "
            "'-' in Item = no number)").font = Font(italic=True, color="777777")
    for j, name in enumerate(COLUMNS, 1):
        c = ws.cell(hr, j, name)
        c.font = bold
        c.fill = fill
        c.border = box
        c.alignment = Alignment(horizontal="center")
    rows = sample.get("items", [])
    for i in range(max(len(rows), 25)):
        row = rows[i] if i < len(rows) else [None] * len(COLUMNS)
        for j, v in enumerate(row, 1):
            c = ws.cell(hr + 1 + i, j, v)
            c.border = box
            if j in (6, 7):
                c.number_format = "#,##0.00"
            if j == 5:
                c.number_format = "00"
    for col, w in zip("ABCDEFG", (16, 44, 22, 8, 8, 14, 14)):
        ws.column_dimensions[col].width = w
    ws["B3"].alignment = Alignment(horizontal="left")
    wb.save(path)


def read_workbook(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Invoice"] if "Invoice" in wb.sheetnames else wb.worksheets[0]
    fields, header_row, cell_of = {}, None, {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        key = str(label).strip().lower() if label is not None else ""
        for name in FIELDS:
            if key == name.lower():
                fields[name] = ws.cell(r, 2).value
                cell_of[name] = (r, 2)
        if (str(ws.cell(r, 2).value or "").strip().lower() == "description"
                and key == "item"):
            header_row = r
            break
    if header_row is None:
        sys.exit("Could not find the Item / Description header row.")

    items = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, j).value for j in range(1, len(COLUMNS) + 1)]
        vals = [v.strip() if isinstance(v, str) else v for v in vals]
        vals = [None if v == "" else v for v in vals]
        if all(v is None for v in vals):
            continue
        items.append(dict(zip(COLUMNS, vals)))
    return fields, items, cell_of


def store_invoice_no(path, cell, number):
    """Write the assigned number back so re-running reuses it."""
    try:
        wb = load_workbook(path)
        ws = wb["Invoice"] if "Invoice" in wb.sheetnames else wb.worksheets[0]
        ws.cell(*cell, number)
        wb.save(path)
    except PermissionError:
        print(f"Note: close {os.path.basename(path)} in Excel and type "
              f"{number} into Invoice No, so a re-run keeps the same number.")


# ---------------------------------------------------------------- lines
def build_lines(items):
    """Turn sheet rows into table lines and work out numbering and totals."""
    lines, counter, total = [], 0, Decimal(0)
    for it in items:
        desc, remark = it["Description"], it["Remark"]
        qty, rate, amount = (to_number(it["Qty"]), to_number(it["Rate"]),
                             to_number(it["Amount"]))
        others = [it["Item"], remark, it["Unit"], qty, rate, amount]
        if desc and all(v is None for v in others):
            lines.append({"heading": str(desc)})
            counter = 0
            continue
        if amount is None and qty is not None and rate is not None:
            amount = qty * rate
        item_no = it["Item"]
        if item_no is None:
            if qty is not None and rate is not None:
                counter += 1
                item_no = counter
        elif str(item_no) == "-":
            item_no = None
        else:
            try:
                counter = int(item_no)
            except (TypeError, ValueError):
                pass
        if amount is not None:
            total += amount
        lines.append({
            "item": "" if item_no is None else str(item_no),
            "desc": "" if desc is None else str(desc),
            "remark": "" if remark is None else str(remark),
            "unit": it["Unit"] or "",
            "qty": "" if it["Qty"] is None else qty_text(it["Qty"]),
            "rate": "" if rate is None else money(rate),
            "amount": "" if amount is None else money(amount),
        })
    return lines, total


# ---------------------------------------------------------------- PDF
def para(text, size=9, bold=False, align=None, leading=None):
    style = ParagraphStyle(
        "p", fontName="Body-Bold" if bold else "Body", fontSize=size,
        leading=leading or size * 1.2,
        alignment={"center": TA_CENTER, "right": TA_RIGHT}.get(align, 0))
    return Paragraph(text, style)


def first_page(canvas, doc, info):
    c = canvas
    w, h = letter
    c.saveState()
    c.drawImage(os.path.join(ASSETS, "logo.png"), 129, h - 143,
                width=102, height=99, mask="auto")
    c.setFont("Body", 11)
    c.drawString(269.3, h - 83, COMPANY["address"])
    c.drawString(268.7, h - 104.5, "Email: ")
    ex = 268.7 + pdfmetrics.stringWidth("Email: ", "Body", 11)
    c.setFillColor(colors.HexColor("#0563C1"))
    c.drawString(ex, h - 104.5, COMPANY["email"])
    c.setLineWidth(0.6)
    c.setStrokeColor(colors.HexColor("#0563C1"))
    c.line(ex, h - 106, ex + pdfmetrics.stringWidth(COMPANY["email"], "Body", 11),
           h - 106)
    c.setFillColor(colors.black)
    c.drawString(270.3, h - 126, f"Hotline: {COMPANY['hotline']}")
    c.drawString(381.7, h - 126, f"Tel: {COMPANY['tel']}")
    # double rule under the letterhead
    c.setStrokeColor(colors.black)
    c.setLineWidth(1.2)
    c.line(61, h - 142, 581, h - 142)
    c.setLineWidth(0.4)
    c.line(100, h - 144.5, 545, h - 144.5)

    c.setFont("Body", 11)
    c.drawString(72, h - 164, "Po No")
    c.drawString(114, h - 164, f"-: {info['po']}")
    c.drawString(72, h - 179, f"Quote No-: {info['quote']}")
    c.setFont("Body", 12)
    c.drawString(72, h - 195, "Date")
    c.drawString(114, h - 195, f"-: {info['date']}")
    c.setFont("Body-Bold", 14)
    c.drawString(352, h - 164, "INVOICE")
    c.setFont("Body-Bold", 12)
    c.drawString(352.3, h - 195, f"Invoice No :{info['number']}")
    c.setFont("Body", 12)
    y = h - 233
    for line in info["customer"]:
        c.drawString(74.8, y, line)
        y -= 20.7
    c.restoreState()


def build_pdf(path, info, lines, total):
    cust_lines = len(info["customer"])
    top = 233 + 20.7 * max(cust_lines - 1, 0) + 24  # below customer block
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=78,
                            rightMargin=50, topMargin=72, bottomMargin=60,
                            title=f"INVOICE {info['number']}",
                            author="Fire Fighting Solutions (Pvt) Ltd")

    widths = [30, 222, 30, 34, 66, 102]  # 484pt, x 78 to 562
    data = [[para(t, 8.5, True, "center") for t in
             ("ITEM", "DESCRIPTION", "UNIT", "QTY", "RATE", "VALUE")]]
    style = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#404040")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for ln in lines:
        if "heading" in ln:
            data.append(["", para(escape(ln["heading"]), 9, True), "", "", "", ""])
            continue
        desc = escape(ln["desc"]).replace("\n", "<br/>")
        if ln["remark"]:
            remark = escape(ln["remark"]).replace("\n", "<br/>")
            desc += ("<br/>" if desc else "") + f"<b>{remark}</b>"
        data.append([para(ln["item"], 9, align="center"), para(desc),
                     para(ln["unit"], 9, align="center"),
                     para(ln["qty"], 9, align="center"),
                     para(ln["rate"], 9, align="right"),
                     para(ln["amount"], 9, align="right")])
    g = len(data)
    data.append(["", para("Grand Total", 11, True, "center"), "", "", "",
                 para(money(total), 11, True, "right")])
    style += [
        ("SPAN", (1, g), (2, g)),
        ("VALIGN", (1, g), (2, g), "BOTTOM"),
        ("VALIGN", (5, g), (5, g), "TOP"),
        ("TOPPADDING", (0, g), (-1, g), 3),
        ("BOTTOMPADDING", (0, g), (-1, g), 4),
    ]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT",
                  minRowHeights=[0] + [20] * (g - 1) + [35])
    table.setStyle(TableStyle(style))

    bank = Table([[para(k, 11), para(f": {v}", 11)] for k, v in COMPANY["bank"]],
                 colWidths=[110, 300], hAlign="LEFT")
    bank.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("TOPPADDING", (0, 0), (-1, -1), 0.5),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5)]))
    sign = Table([[para("……………………………………..", 11, True),
                   para("………………………………………..", 11, True)],
                  [para("Prepared By", 11), para("Authorized Signature", 11)]],
                 colWidths=[304, 180], hAlign="LEFT")
    sign.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("LEFTPADDING", (0, 1), (-1, 1), 14),
                              ("TOPPADDING", (0, 1), (-1, 1), 8)]))

    indent = Table([[x] for x in (
        para("<u>Remittance Information</u>", 11), Spacer(1, 12), bank)],
        hAlign="LEFT")
    indent.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 1),
                                ("TOPPADDING", (0, 0), (-1, -1), 0),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))

    story = [Spacer(1, top - 72), table, Spacer(1, 40),
             KeepTogether([indent, Spacer(1, 80), sign])]
    doc.build(story, onFirstPage=lambda c, d: first_page(c, d, info))


# ---------------------------------------------------------------- main
def format_date(value):
    if value is None:
        return dt.date.today().strftime("%d/%m/%Y")
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%d/%m/%Y")
    return str(value)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("workbook", nargs="?", help="filled-in invoice .xlsx")
    ap.add_argument("--template", metavar="XLSX", help="write a blank workbook")
    ap.add_argument("--next", action="store_true", help="show next number")
    ap.add_argument("-o", "--out", help="output PDF path")
    args = ap.parse_args()

    if args.template:
        write_template(args.template)
        print(f"Blank invoice workbook written: {args.template}")
        return
    if args.next:
        print(read_counter() + 1)
        return
    if not args.workbook:
        ap.print_help()
        return

    register_fonts()
    fields, items, cell_of = read_workbook(args.workbook)
    lines, total = build_lines(items)
    if not lines:
        sys.exit("No items found in the workbook.")

    last = read_counter()
    number = fields.get("Invoice No")
    assigned = number is None
    if assigned:
        number = last + 1
    number = int(number) if str(number).isdigit() else number

    customer = [str(fields[k]) for k in
                ("Customer", "Address line 1", "Address line 2", "Address line 3")
                if fields.get(k) not in (None, "")]
    info = {"number": number, "date": format_date(fields.get("Date")),
            "po": fields.get("PO No") or "", "quote": fields.get("Quote No") or "",
            "customer": customer}

    os.makedirs(OUTPUT, exist_ok=True)
    out = args.out or os.path.join(OUTPUT, f"INVOICE {number}.pdf")
    build_pdf(out, info, lines, total)
    if isinstance(number, int) and number > last:
        write_counter(number)
    if assigned and "Invoice No" in cell_of:
        store_invoice_no(args.workbook, cell_of["Invoice No"], number)
    print(f"Invoice {number}  total {money(total)}  ->  {out}")


if __name__ == "__main__":
    main()
